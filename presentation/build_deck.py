"""Build the content-only Reach deck and its reviewable Markdown companion.

Narrative content lives in story_slides.py; technical appendix content is below.
Install presentation/requirements.txt, then run from the repository root:
  python presentation/build_deck.py
"""
from pathlib import Path
import argparse
import sys

ROOT = Path(__file__).resolve().parents[1]
LOCAL_DEPS = ROOT / '.tmp-presentation-deps'
if LOCAL_DEPS.exists():
    sys.path.insert(0, str(LOCAL_DEPS))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path(__file__).resolve().parent
VERSION = '1.0'
OUTPUT_STEM = f'Reach-developer-presentation-v{VERSION}'
from story_slides import MAIN_SLIDES, BENEFITS_SLIDE

SLIDES = MAIN_SLIDES + [BENEFITS_SLIDE,
    {'title': 'Appendix: getting started',
 'subtitle': 'Install once, then adopt a repository.',
 'appendix': True,
 'bullets': ['Requires Claude Code and PowerShell. Objects-mode landing needs Git 2.38 or newer.',
             'claude plugin marketplace add SpoiledInk13/Reach',
             'claude plugin install reach@reach',
             'In Claude Code: /reach:adopt',
             'Then use /reach:ideate, /reach:build, and /reach:milestone.'],
 'notes': 'These are the commands documented in this repository. Use Windows PowerShell 5.1 or pwsh; Git is '
          'required in either landing mode. Objects-mode landing needs merge-tree --write-tree from Git '
          "2.38. Adoption is a conversation about the repository's existing workflow, not a blind "
          'scaffolding step. It writes process.json and the script entry point, configures verification, and '
          'ignores local logs and locks. On an already-adopted repository, it audits for gaps instead of '
          'repeating the original replacement. Milestone is optional when there is no outcome that needs '
          'human judgment. Commands were sourced locally; external compatibility was not rechecked for the '
          'presentation.',
 'source': 'README.md: How to use it; scripts/Land.ps1; scripts/lib/Common.ps1: Get-PowerShellExe; '
           'skills/adopt/SKILL.md'},
    {'title': 'Appendix: where the workflow lives',
 'subtitle': 'The decisions and checks live with the repository.',
 'appendix': True,
 'bullets': ['process.json: document locations, size limits, evidence conventions, checks, and lanes.',
             "Architecture document: the system's structure, shared rules, and list of units.",
             'Unit documents: responsibilities, behavior, dependencies, and evidence.',
             'Walkthroughs: the experiences that still need human acceptance.',
             'Scripts/reach.ps1: the entry point for verification and lane operations.'],
 'notes': 'The architecture document is called the spine in Reach. Unit documents describe smaller parts of '
          'the system. Their state is built or unbuilt; a new claim on a built unit can carry owed evidence '
          'until implemented. The gate checks the links between evidence and documents, while tests check '
          "the behavior. The script entry point resolves the installed plugin so the repository doesn't "
          "hard-code a versioned cache path. The template's reach version, 0.0.0, is a placeholder; adoption "
          'records the actual baseline. Audit compares that baseline with the running plugin in both '
          'directions and reports which plugin path answered.',
 'source': 'templates/process.json; templates/unit.md; README.md: process.json'},
    {'title': 'Appendix: checking the process',
 'subtitle': 'Use the same entry point locally and in automation.',
 'appendix': True,
 'bullets': ['pwsh Scripts/reach.ps1 gate — check repository constraints and evidence links.',
             'pwsh Scripts/reach.ps1 all — run the gate and automated verification tiers.',
             'pwsh Scripts/reach.ps1 prove — exercise the 50 supplied negative controls.',
             'pwsh Scripts/reach.ps1 audit — find adoption gaps and version mismatches.',
             'Add custom checks when a real defect shows the need, then prove they can fail.'],
 'notes': 'A custom check is a PowerShell file in the configured checks directory with a matching '
          "Test-<filename> function. The supplied proof suite doesn't automatically cover new custom checks "
          'or application behavior. A recent archive-check fix illustrates the need for realistic controls: '
          'it caught edits before commit but missed them after commit. It now checks both uncommitted '
          'changes and branch commits, including deletion of the last archive file. Controls can also check '
          'diagnostic wording. An unset document cap is now reported as missing configuration, and audit '
          'identifies version drift. Optional unattended operation requires explicit opt-in and disables '
          'agent permission prompts. Its supervisor refuses a namespaced command if the plugin is absent '
          'from the installed-plugin registry; that is a presence check, not full availability validation.',
 'source': 'README.md: Verify, Writing your own checks, Unattended; scripts/Audit.ps1; scripts/Run-Lane.ps1'},
]


def textbox(slide, x, y, w, h, text, size, bold=False, color=(0, 0, 0)):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = 'Arial'
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(*color)
    return shape


def build(output_name=f'{OUTPUT_STEM}.pptx'):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = f'Reach — developer presentation v{VERSION}'
    prs.core_properties.subject = 'Approved content and flow; visual style and assets deferred'
    prs.core_properties.version = VERSION
    prs.core_properties.author = 'Amelia Bleeker'
    main_count = sum(not item.get('appendix', False) for item in SLIDES)
    appendix_count = len(SLIDES) - main_count
    markdown = [f'# Reach — developer presentation v{VERSION}',
                f'{main_count} main slides + {appendix_count} appendix slides. Approximately 20-25 minutes plus discussion. '
                'Plain editable layout; visual style and assets intentionally deferred.',
                'Pip and the cottage fireplace are a fictional teaching story. Speaker notes distinguish repository mechanisms, presenter observations, '
                'inferred benefits, and fictional events. Developer remarks in subtitles are illustrative, not research quotations. '
                'The presentation contains no live demonstration or claimed application test results.']
    for i, item in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        textbox(slide, .65, .38, 12, .78, item['title'], 30, True)
        textbox(slide, .65, 1.3, 12, .8, item['subtitle'], 21, color=(65, 65, 65))
        count = len(item['bullets'])
        row_height = min(1.02, 4.5 / count)
        for j, line in enumerate(item['bullets']):
            textbox(slide, .8, 2.22 + j * row_height, 11.75, row_height - .08,
                    '\u2022 ' + line, 22)
        if item.get('payoff'):
            textbox(slide, .8, 6.55, 11.75, .4, item['payoff'], 17, True)
        label = f'v{VERSION} | Appendix' if item.get('appendix') else f'v{VERSION}'
        textbox(slide, .65, 7.06, 12, .25, f'Reach | {label} | {i:02d}', 10, color=(95, 95, 95))
        slide.notes_slide.notes_text_frame.text = (
            f"SLIDE {i}: {item['title']}\n\n{item['notes']}"
            + (f"\n\nPRACTICAL BENEFIT: {item['payoff']}" if item.get('payoff') else '')
            + f"\n\nSOURCE BASIS: {item['source']}"
        )
        markdown.extend([f"## {i:02d}. {item['title']}", item['subtitle'],
                         '\n'.join('- ' + x for x in item['bullets'])])
        if item.get('payoff'):
            markdown.append('**Practical benefit:** ' + item['payoff'])
        markdown.extend(['**Speaker notes**\n\n' + item['notes'],
                         '**Source basis:** ' + item['source']])
    target = OUT / output_name
    prs.save(target)
    (OUT / f'{OUTPUT_STEM}.md').write_text('\n\n'.join(markdown) + '\n', encoding='utf-8')
    # Check the saved package, notes, and geometry, rather than only the in-memory deck.
    saved = Presentation(target)
    assert len(saved.slides) == len(SLIDES)
    for slide, item in zip(saved.slides, SLIDES):
        assert item['notes'] in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= saved.slide_width
            assert shape.top + shape.height <= saved.slide_height
        texts = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        assert item['title'] in texts
        assert all(line in texts for line in item['bullets'])
        if item.get('payoff'):
            assert item['payoff'] in texts
    print(f'Created and validated {len(saved.slides)} editable slides with speaker notes: {target}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default=f'{OUTPUT_STEM}.pptx',
                        help='Output filename inside presentation/')
    build(parser.parse_args().output)
