"""Build the illustrated edition using the approved v1.0 slide copy.

python presentation/build_storybook.py
"""
from pathlib import Path
from build_deck import SLIDES
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont

ROOT = Path(__file__).resolve().parent
ART = ROOT / 'artwork'
OUTPUT = ROOT / 'Reach-developer-presentation-storybook.pptx'
PAPER = 'F7F2E7'
INK = '253D35'
MUTED = '667166'
RUST = 'A44F32'
GOLD = '9B742D'
LINE = 'D8CEB9'
CARD = 'EDE6D7'
WHITE = 'FFFCF5'
GREEN = '3D7258'
FONT = 'Calibri'
TITLE = 'Georgia'

ROLES = [
    ('REACH', INK), ('THE ORIGIN', INK), ('THE STORY', INK), ('THE DESIGN GAP', RUST),
    ('THE WORKFLOW', INK), ('IDEATE', INK), ('BUILD', RUST), ('BUILD', RUST),
    ('IDEATE', INK), ('IDEATE', INK), ('IDEATE → BUILD', INK), ('BUILD', RUST),
    ('REACH SCRIPTS', INK), ('MILESTONE', GOLD), ('THE HANDOFF', INK),
    ('IN PRACTICE', INK), ('PROJECT FIT', INK), ('BENEFITS', INK),
    ('ADOPT', INK), ('REPOSITORY MAP', INK), ('VERIFICATION', RUST),
]
SCENES = {
    1: ('pip-paperwork.png', 'The process grew. The playable goal stayed small.'),
    2: ('pip-cottage-open.png', 'One destination: the warm fireplace inside.'),
    3: ('pip-door-closed.png', 'The route is blocked. The design decision is still open.'),
    5: ('pip-cottage-open.png', 'Decide what the player should expect.'),
    7: ('pip-door-closed.png', 'A blocked claim becomes a question for ideate.'),
    9: ('pip-window-open.png', 'Same destination. A second kind of opening.'),
    13: ('pip-fireplace.png', 'Working behavior still needs a human verdict.'),
    14: ('pip-fireplace.png', 'The warmth is the outcome. The contract carries the understanding.'),
    15: ('pip-cottage-open.png', 'Think together about the whole game.'),
}


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = .07
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape


def text(slide, x, y, w, h, value, size=20, color=INK, bold=False, font=FONT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.08
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape._element.spPr.append(OxmlElement('a:effectLst'))


def card(slide, x, y, w, h, heading, body, accent=INK):
    box(slide, x, y, w, h, WHITE, LINE, True)
    box(slide, x, y, .055, h, accent)
    text(slide, x+.2, y+.15, w-.4, .36, heading, 19, accent, True)
    text(slide, x+.2, y+.64, w-.4, h-.72, body, 18)


def picture(slide, name, x, y, w):
    # Preserve the full composition and aspect ratio; never crop Pip or the entrances.
    return slide.shapes.add_picture(str(ART/name), Inches(x), Inches(y), width=Inches(w))


def diagram(slide, index):
    x, w = 7.5, 5.15
    if index == 6:
        for i, (head, body, color) in enumerate([
            ('01  AGREED BEHAVIOR', 'Closed door → Pip stays put', INK),
            ('02  NAMED EVIDENCE', 'Position + blocked-route feedback', RUST),
            ('03  VERIFICATION', 'Tests → game build → explicit result', GREEN),
        ]):
            card(slide,x,2.22+i*1.15,w,1.02,head,body,color)
    elif index == 8:
        card(slide,x,2.25,w,1.15,'BUILD → OPEN QUESTION','What if the door closes mid-walk?',RUST)
        text(slide,x+2.3,3.48,.6,.4,'↓',24,MUTED)
        card(slide,x,4.02,w,1.42,'IDEATE → UPDATED CONTRACT','Stop before the obstacle.\nKeep the last valid position.',INK)
        text(slide,x,5.66,w,.42,'Build reads the answer from the document.',17,MUTED)
    elif index == 10:
        card(slide,x,2.22,w,1.46,'OPENING','Open / closed\nLocked / unlocked',INK)
        line(slide,x+w/2,3.68,x+w/2,3.98)
        line(slide,x+1.2,3.98,x+w-1.2,3.98)
        line(slide,x+1.2,3.98,x+1.2,4.16)
        line(slide,x+w-1.2,3.98,x+w-1.2,4.16)
        card(slide,x,4.16,2.45,1.2,'DOOR','Shared rules',RUST)
        card(slide,x+2.7,4.16,2.45,1.2,'WINDOW','Shared rules',RUST)
        text(slide,x,5.59,w,.62,'Traverse when open, reachable, and large enough.',18,MUTED)
    elif index == 11:
        rows=[('BREAK IT','Closed opening allows passage',RUST),('EXPECT RED','The movement assertion fails',RUST),('RESTORE','Door + window cases pass',GREEN),('KEEP RUNNING','Catch the rule breaking again',GREEN)]
        for i,(h,b,c) in enumerate(rows):
            box(slide,x,2.15+i*.94,w,.8,WHITE,LINE,True)
            text(slide,x+.18,2.27+i*.94,1.58,.45,h,15,c,True)
            text(slide,x+1.83,2.27+i*.94,3.08,.52,b,17)
    elif index == 12:
        card(slide,x,2.22,2.45,1.22,'BUILD','Own worktree',RUST)
        card(slide,x+2.7,2.22,2.45,1.22,'MILESTONE','Own worktree',GOLD)
        text(slide,x,3.67,w,.4,'Verified commit → checked merge',19,INK,True)
        card(slide,x,4.3,w,1.18,'INTEGRATION','Publish the accepted decisions and code.',INK)
        text(slide,x,5.72,w,.42,'Ideate syncs before the next conversation.',17,MUTED)


def fit_check(prs):
    issues=[]
    for i,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left+sh.width <= prs.slide_width+10
            assert sh.top+sh.height <= prs.slide_height+10
            if not sh.has_text_frame or not sh.text:
                continue
            available_w=(sh.width-sh.text_frame.margin_left-sh.text_frame.margin_right)/914400*96
            available_h=(sh.height-sh.text_frame.margin_top-sh.text_frame.margin_bottom)/914400*96
            needed=0
            for p in sh.text_frame.paragraphs:
                size=p.font.size.pt
                filename=('georgiab.ttf' if p.font.bold else 'georgia.ttf') if p.font.name==TITLE else ('calibrib.ttf' if p.font.bold else 'calibri.ttf')
                font=ImageFont.truetype('C:/Windows/Fonts/'+filename,round(size*96/72))
                count=1
                current=''
                for word in p.text.split():
                    trial=(current+' '+word).strip()
                    if current and font.getlength(trial)>available_w:
                        count+=1
                        current=word
                    else:
                        current=trial
                needed+=count*size*96/72*1.15
            if needed>available_h+1:
                issues.append((i,sh.text,round(needed),round(available_h)))
    assert not issues,issues


def build():
    prs=Presentation()
    prs.slide_width,prs.slide_height=Inches(13.333), Inches(7.5)
    prs.core_properties.title='Reach — storybook visual edition'
    prs.core_properties.subject='Approved v1.0 content with storybook artwork and editable diagrams'
    prs.core_properties.author='Amelia Bleeker'
    for index,item in enumerate(SLIDES):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb=rgb(PAPER)
        role,accent=ROLES[index]
        box(s,.58,.43,.11,.18,accent)
        text(s,.83,.39,11,.3,role if index==0 else role+'  /  REACH',12,accent,True)
        if index==0:
            box(s,6.6,0,6.733,7.5,INK)
            picture(s,'pip-cottage-open.png',6.82,1.5,6.28)
            text(s,.7,1.27,5.55,2.17,item['title'],35,INK,False,TITLE)
            text(s,.7,3.58,5.3,.85,item['subtitle'],22,MUTED)
            for j,b in enumerate(item['bullets']):
                text(s,.7,4.65+j*.84,5.25,.75,b,20)
            text(s,7.04,6.06,5.5,.76,'A story about intent, implementation,\nand knowing what is true.',20,PAPER)
        else:
            text(s,.65,.95,12.05,.66,item['title'],28,INK,False,TITLE)
            text(s,.68,1.7,12,.43,item['subtitle'],19,MUTED)
            line(s,.68,2.03,12.65,2.03)
            if index==4:
                for j,(h,b,c) in enumerate(zip(['IDEATE','BUILD','MILESTONE'],item['bullets'],[INK,RUST,GOLD])):
                    x=.68+j*4.06
                    card(s,x,2.64,3.82,2.62,h,b,c)
                    text(s,x+.2,5.49,3.4,.5,['Your judgment','Automated evidence','Your experience'][j],20,c,True)
            else:
                split=index in SCENES or index in (6,8,10,11,12)
                w=6.28 if split else 11.52
                row=.94 if len(item['bullets'])<=4 else .86
                for j,b in enumerate(item['bullets']):
                    text(s,.72,2.37+j*row,.34,.36,f'{j+1:02}',11,accent,True)
                    text(s,1.16,2.3+j*row,w-.48,row-.09,b,20 if split else 22)
                if index in SCENES:
                    name,caption=SCENES[index]
                    picture(s,name,7.35,2.3,5.3)
                    text(s,7.38,5.97,5.25,.57,caption,16,MUTED)
                elif split:
                    diagram(s,index)
            if item.get('payoff'):
                box(s,.66,6.66,12,.06,accent)
                text(s,.72,6.84,11.95,.34,item['payoff'],16,accent,True)
        text(s,.68,7.24,11.55,.18,'REACH  ·  CONTENT v1.0  ·  STORYBOOK EDITION'+('  /  APPENDIX' if item.get('appendix') else ''),9,MUTED)
        text(s,12.12,7.2,.5,.23,f'{index+1:02}',11,MUTED)
        s.notes_slide.notes_text_frame.text=(f"SLIDE {index+1}: {item['title']}\n\n{item['notes']}"
            +(f"\n\nPRACTICAL BENEFIT: {item['payoff']}" if item.get('payoff') else '')
            +f"\n\nSOURCE BASIS: {item['source']}\n\nART: AI-generated fictional storybook scene; diagrams are editable PowerPoint objects.")
    fit_check(prs)
    prs.save(OUTPUT)
    check=Presentation(OUTPUT)
    for s,item in zip(check.slides,SLIDES):
        text_content='\n'.join(sh.text for sh in s.shapes if sh.has_text_frame)
        assert item['title'] in text_content
        assert item['subtitle'] in text_content
        assert all(b in text_content for b in item['bullets'])
        assert item['notes'] in s.notes_slide.notes_text_frame.text
    print(f'Built {OUTPUT}; {len(check.slides)} slides; approved copy preserved; text fit and geometry checked.')


if __name__=='__main__':
    build()
